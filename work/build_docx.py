#!/usr/bin/env python3
"""Build the Chinese Word document from content.json + work/trans/*.json.

Translations: work/trans/*.json  ->  {"<content-index>": "中文", ...}
Missing translations fall back to the German text (marked) so the layout can be
validated before the translation is complete.
"""
import json, os, struct, glob
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

WORK = os.path.dirname(__file__)
ROOT = os.path.dirname(WORK) or '.'
CONTENT_W = 6.1                      # inches, usable text width
ZOOM = 4
CJK = '宋体'
CJK_H = '黑体'                       # headings

def png_px(path):
    with open(path, 'rb') as f:
        head = f.read(24)
    return struct.unpack('>I', head[16:20])[0], struct.unpack('>I', head[20:24])[0]

def img_inches(path):
    w, h = png_px(path)
    return w / ZOOM / 72.0, h / ZOOM / 72.0

def set_cjk(run, font, size=None, bold=None, color=None):
    run.font.name = font
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), font)
    rFonts.set(qn('w:ascii'), font)
    rFonts.set(qn('w:hAnsi'), font)
    if size is not None: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if color is not None: run.font.color.rgb = RGBColor(*color)

def load_translations():
    tr = {}
    for fp in sorted(glob.glob(os.path.join(WORK, 'trans', '*.json'))):
        with open(fp) as f:
            for k, v in json.load(f).items():
                tr[int(k)] = v
    return tr

# images recovered by hand for captions the extractor missed (sub-figure /
# inline captions). Inserted AFTER the given content index so existing indices
# (and thus the translation mapping) are preserved.
EXTRA_AFTER = {
    76: [("img/figure_2_9.png", None)],                       # Fig 2.9 (caption = idx 77)
    347: [("img/figure_4_11.png", ("4.11", "IQ 相位评估精度，61 GHz FMCW 雷达 [71] © 2014 IEEE。"))],
}

def main():
    content = json.load(open(os.path.join(WORK, 'content.json')))
    tr = load_translations()

    doc = Document()
    # base style
    normal = doc.styles['Normal']
    normal.font.size = Pt(10.5)
    normal.font.name = CJK
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), CJK)
    for sec in doc.sections:
        sec.left_margin = Inches(1.0); sec.right_margin = Inches(1.0)
        sec.top_margin = Inches(0.9); sec.bottom_margin = Inches(0.9)

    # ---------- title page ----------
    def tline(text, size, bold=False, font=CJK_H, after=6, color=None):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(after)
        r = p.add_run(text); set_cjk(r, font, size, bold, color)
        return p
    for _ in range(2): doc.add_paragraph()
    tline('用于高精度距离测量的', 22, True)
    tline('FMCW 雷达信号处理', 22, True, after=20)
    tline('FMCW-Radarsignalverarbeitung zur Entfernungsmessung mit hoher Genauigkeit',
          12, False, CJK, after=28, color=(0x55, 0x55, 0x55))
    tline('作者：Steffen Scherr', 13, False, after=6)
    tline('卡尔斯鲁厄理工学院（KIT）高频技术与电子学研究所', 12, False, after=6)
    tline('卡尔斯鲁厄高频技术与电子学研究所研究报告  第 83 卷', 11, False,
          CJK, after=24, color=(0x55, 0x55, 0x55))
    tline('（中文译本）', 12, False, after=6)
    doc.add_page_break()

    # ---------- body ----------
    def add_image(img_path, max_w=CONTENT_W):
        full = os.path.join(WORK, img_path)
        if not os.path.exists(full):
            raise FileNotFoundError(full)
        w_in, h_in = img_inches(full)
        if w_in > max_w:
            w_in = max_w
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(4)
        p.add_run().add_picture(full, width=Inches(w_in))

    def add_caption(prefix, zh):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run(prefix); set_cjk(r, CJK, 9, True)
        r2 = p.add_run(zh); set_cjk(r2, CJK, 9, False)

    miss = 0
    for i, e in enumerate(content):
        t = e['type']
        if t == 'heading':
            zh = tr.get(i, '〖' + e['text'] + '〗')
            if i not in tr: miss += 1
            lvl = e.get('level', 2)
            num = e.get('num', '')
            if lvl == 1:
                doc.add_page_break()
                p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(16)
                r = p.add_run(f'第 {num} 章  ' if num else ''); set_cjk(r, CJK_H, 20, True, (0x1F, 0x3A, 0x5F))
                r = p.add_run(zh); set_cjk(r, CJK_H, 20, True, (0x1F, 0x3A, 0x5F))
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(12 if lvl == 2 else 8)
                p.paragraph_format.space_after = Pt(6)
                sz = 15 if lvl == 2 else 12.5
                r = p.add_run((num + '  ') if num else ''); set_cjk(r, CJK_H, sz, True, (0x1F, 0x3A, 0x5F))
                r = p.add_run(zh); set_cjk(r, CJK_H, sz, True, (0x1F, 0x3A, 0x5F))
        elif t == 'para':
            zh = tr.get(i)
            if zh is None: zh = '〖待译〗' + e['text']; miss += 1
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.first_line_indent = Inches(0.28)
            r = p.add_run(zh); set_cjk(r, CJK, 10.5)
        elif t == 'equation':
            add_image(e['img'])
        elif t in ('figure', 'table'):
            if e.get('img'):
                add_image(e['img'])
            zh = tr.get(i, e.get('caption', ''))
            if e.get('caption') and i not in tr: miss += 1
            label = '图' if t == 'figure' else '表'
            add_caption(f"{label} {e.get('num','')}：", zh)
        # inject hand-recovered images whose captions the extractor missed
        for img_rel, cap in EXTRA_AFTER.get(i, []):
            add_image(img_rel)
            if cap:
                add_caption(f"图 {cap[0]}：", cap[1])

    out = os.path.join(ROOT, 'FMCW雷达信号处理_中文译本.docx')
    doc.save(out)
    print('Saved:', out)
    print('Missing translations:', miss, '/ elements:', len(content))

if __name__ == '__main__':
    main()
