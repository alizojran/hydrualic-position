#!/usr/bin/env python3
"""Build the Chinese Word documents for the two radar papers.

content_<doc>.json (structure + English) + trans_<doc>.json (index -> 中文)
                                          -> ../<出力>.docx
A translation of "" removes that element (used to drop caption-leak fragments).
References and the page-0 footnote are kept verbatim (English) by convention.
"""
import json, os, struct
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

WORK = os.path.dirname(__file__)
ROOT = os.path.dirname(WORK) or '.'
CONTENT_W = 6.3
ZOOM = 4
CJK = '宋体'
CJK_H = '黑体'
NAVY = (0x1F, 0x3A, 0x5F)
GRAY = (0x55, 0x55, 0x55)

DOCS = {
    'journal': dict(
        title_en='High-Accuracy Range Detection Radar Sensor for Hydraulic Cylinders',
        source='IEEE Sensors Journal, 第 14 卷, 第 3 期, 2014 年 3 月, 第 734–745 页',
        out='High-Accuracy_Range_Detection_Radar_Sensor_中文译本.docx'),
    'conf': dict(
        title_en='FMCW Radar in Oil-filled Waveguides for Range Detection in Hydraulic Cylinders',
        source='Proceedings of the 9th European Radar Conference (EuRAD), 2012 年 10 月 31 日–11 月 2 日, '
               '荷兰阿姆斯特丹, 第 63–66 页',
        out='C2012_Ayhan_FMCW_oil-filled_waveguides_中文译本.docx'),
}


def png_px(path):
    with open(path, 'rb') as f:
        head = f.read(24)
    return struct.unpack('>I', head[16:20])[0], struct.unpack('>I', head[20:24])[0]


def img_inches(path):
    w, h = png_px(path)
    return w / ZOOM / 72.0, h / ZOOM / 72.0


def set_cjk(run, font, size=None, bold=None, color=None, italic=None):
    run.font.name = font
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    for a in ('w:eastAsia', 'w:ascii', 'w:hAnsi'):
        rFonts.set(qn(a), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def build(doc_key):
    cfg = DOCS[doc_key]
    content = json.load(open(os.path.join(WORK, f'content_{doc_key}.json')))
    trans = {}
    tp = os.path.join(WORK, f'trans_{doc_key}.json')
    if os.path.exists(tp):
        trans = {int(k): v for k, v in json.load(open(tp)).items()}

    doc = Document()
    normal = doc.styles['Normal']
    normal.font.size = Pt(10.5)
    normal.font.name = CJK
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), CJK)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(0.95)
        sec.top_margin = sec.bottom_margin = Inches(0.9)

    def para(space_after=6, align=None, before=None, indent=None, lh=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        if before is not None:
            p.paragraph_format.space_before = Pt(before)
        if align is not None:
            p.alignment = align
        if indent is not None:
            p.paragraph_format.first_line_indent = Inches(indent)
        if lh is not None:
            p.paragraph_format.line_spacing = lh
        return p

    def add_image(img_rel, max_w=CONTENT_W):
        full = os.path.join(WORK, img_rel)
        if not os.path.exists(full):
            return
        w_in, h_in = img_inches(full)
        if w_in > max_w:
            w_in = max_w
        p = para(4, WD_ALIGN_PARAGRAPH.CENTER, before=4)
        p.add_run().add_picture(full, width=Inches(w_in))

    def caption(prefix, zh):
        p = para(10, WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(prefix); set_cjk(r, CJK, 9, True, NAVY)
        r2 = p.add_run(zh); set_cjk(r2, CJK, 9, False)

    miss = []

    # find front-matter elements
    fm = {e['type']: i for i, e in enumerate(content) if e['type'] in ('title', 'authors', 'abstract', 'indexterms')}

    # ---------------- title page ----------------
    for _ in range(2):
        doc.add_paragraph()
    ti = fm.get('title')
    title_zh = trans.get(ti, content[ti]['text']) if ti is not None else ''
    p = para(8, WD_ALIGN_PARAGRAPH.CENTER); r = p.add_run(title_zh); set_cjk(r, CJK_H, 22, True, NAVY)
    p = para(26, WD_ALIGN_PARAGRAPH.CENTER); r = p.add_run(cfg['title_en']); set_cjk(r, CJK, 12, False, GRAY, italic=True)
    ai = fm.get('authors')
    if ai is not None:
        for line in content[ai]['text'].split('\n'):
            p = para(2, WD_ALIGN_PARAGRAPH.CENTER); r = p.add_run(line); set_cjk(r, CJK, 11, False)
    p = para(24, WD_ALIGN_PARAGRAPH.CENTER); r = p.add_run(cfg['source']); set_cjk(r, CJK, 10.5, False, GRAY)
    p = para(4, WD_ALIGN_PARAGRAPH.CENTER); r = p.add_run('（中文译本）'); set_cjk(r, CJK, 11, False)
    p = para(0, WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run('本译本由原文翻译生成，公式、图、表自原文裁切保留；参考文献保持原文。')
    set_cjk(r, CJK, 8.5, False, GRAY)
    doc.add_page_break()

    # ---------------- body ----------------
    for i, e in enumerate(content):
        t = e['type']
        if t in ('title', 'authors'):
            continue
        if t == 'abstract':
            p = para(6, before=2); r = p.add_run('摘要'); set_cjk(r, CJK_H, 13, True, NAVY)
            zh = trans.get(i);
            if zh is None:
                miss.append(i); zh = '〖待译〗' + e['text']
            p = para(8, indent=0.28, lh=1.25); r = p.add_run(zh); set_cjk(r, CJK, 10.5, italic=True)
        elif t == 'indexterms':
            zh = trans.get(i)
            if zh is None:
                miss.append(i); zh = e['text']
            p = para(14, indent=0.0)
            r = p.add_run('关键词：'); set_cjk(r, CJK_H, 10.5, True, NAVY)
            r2 = p.add_run(zh); set_cjk(r2, CJK, 10.5, italic=True)
        elif t == 'note':
            p = para(10, lh=1.0)
            r = p.add_run(e['text']); set_cjk(r, CJK, 8, False, GRAY, italic=True)
        elif t == 'heading':
            zh = trans.get(i)
            if zh is None:
                miss.append(i); zh = '〖' + e['text'] + '〗'
            num = e.get('num', '')
            if e['level'] == 1:
                p = para(6, before=14)
                if num:
                    r = p.add_run(num + '.  '); set_cjk(r, CJK_H, 15, True, NAVY)
                r = p.add_run(zh); set_cjk(r, CJK_H, 15, True, NAVY)
            else:
                p = para(5, before=9)
                if num:
                    r = p.add_run(num + '.  '); set_cjk(r, CJK_H, 12, True, NAVY)
                r = p.add_run(zh); set_cjk(r, CJK_H, 12, True, NAVY)
        elif t == 'para':
            zh = trans.get(i)
            if zh is None:
                miss.append(i); zh = '〖待译〗' + e['text']
            if zh.strip() == '':
                continue            # dropped (caption-leak fragment / merged elsewhere)
            p = para(6, indent=0.28, lh=1.25); r = p.add_run(zh); set_cjk(r, CJK, 10.5)
        elif t == 'equation':
            add_image(e['img'])
        elif t in ('figure', 'table'):
            if e.get('img'):
                add_image(e['img'])
            zh = trans.get(i)
            if zh is None:
                miss.append(i); zh = e.get('caption', '')
            label = '图' if t == 'figure' else '表'
            caption(f"{label} {e.get('num','')}　", zh)
        elif t == 'references':
            p = para(6, before=14); r = p.add_run('参考文献'); set_cjk(r, CJK_H, 15, True, NAVY)
            for it in e['items']:
                p = para(2, lh=1.05); p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.first_line_indent = Inches(-0.3)
                r = p.add_run(it); set_cjk(r, CJK, 8.5, False)

    out = os.path.join(ROOT, cfg['out'])
    doc.save(out)
    print(f'[{doc_key}] saved {cfg["out"]} | missing translations: {len(miss)} {miss[:12]}')


if __name__ == '__main__':
    for k in ('journal', 'conf'):
        build(k)
